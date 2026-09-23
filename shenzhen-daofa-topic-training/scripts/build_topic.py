#!/usr/bin/env python3
"""深圳道法试卷构建脚本。

用法:
    python build_exam.py --src <试卷.html> --out <输出目录> --title "九上道法第X单元测试卷"

功能:
    1. 将合一版试卷 HTML 拆分为 学生版 / 答案版 两个 HTML
    2. 调用 html_to_docx 插件转为两个 DOCX (A4)
    3. 用 python-docx 自动校验格式 (全黑无彩色 / 答案表格 / 学生版无答案)

退出码: 0 = 全部通过, 1 = 有校验项失败。
"""
import argparse
import glob
import os
import re
import subprocess
import sys

HOME = os.path.expanduser("~")
DOCX_PLUGIN_GLOB = os.path.join(
    HOME, ".workbuddy", "plugins", "cache", "workbuddy-builtin",
    "tencent-docx", "*", "skills", "html-to-docx", "scripts")
HTMLDOCX_LIBS = os.path.join(
    HOME, ".workbuddy", "binaries", "python", "envs", "htmldocx-libs")


def find_html_to_docx():
    """定位 html_to_docx 插件 scripts 目录。"""
    for path in sorted(glob.glob(DOCX_PLUGIN_GLOB), reverse=True):
        if os.path.isfile(os.path.join(path, "html_to_docx")) or \
           os.path.isdir(os.path.join(path, "html_to_docx")):
            return path
    raise SystemExit("[ERROR] 未找到 html_to_docx 插件, 请检查路径: %s" % DOCX_PLUGIN_GLOB)


def _absolutize_images(html, src_path):
    """把 <img src> 相对路径改为绝对路径。

    拆分后的 student/answers HTML 写到了输出目录，相对路径会解析失败导致图片丢失，
    以源 HTML 所在目录为基准转为绝对路径（正斜杠）。
    """
    base = os.path.dirname(os.path.abspath(src_path))

    def repl(m):
        src = m.group(1)
        if os.path.isabs(src) or src.startswith(("http://", "https://")):
            return m.group(0)
        abs_src = os.path.join(base, src).replace("\\", "/")
        return m.group(0).replace('src="%s"' % src, 'src="%s"' % abs_src)

    return re.sub(r'<img[^>]*?src="([^"]+)"', repl, html)


def split_html(src_path, out_dir, stem):
    """把合一版 HTML 拆成 学生版 / 答案版, 返回两个文件路径。"""
    with open(src_path, encoding="utf-8") as f:
        html = f.read()

    html = _absolutize_images(html, src_path)

    m = re.search(r'<div class="answer-key">.*</div>\s*</body>', html, re.S)
    if not m:
        raise SystemExit("[ERROR] 未找到 <div class=\"answer-key\"> 答案块, 请检查 HTML 结构")
    answer_block = m.group(0)

    # 学生版: 移除答案块 (answer_block 末尾含 </body>, 用其补回闭合标签)
    student_html = html.replace(answer_block, "</body>")
    if "answer-key" in re.sub(r"\.answer-key\s*\{[^}]*\}", "", student_html):
        raise SystemExit("[ERROR] 学生版仍残留 answer-key 内容, 请检查 div 嵌套闭合")

    student_path = os.path.join(out_dir, stem + "-student.html")
    with open(student_path, "w", encoding="utf-8") as f:
        f.write(student_html)

    # 答案版: 头部 + 答案块
    head = html.split("<body>")[0] + "<body>\n"
    ans_start = answer_block.index('<p class="exam-head">')
    ans_body = answer_block[ans_start:].replace("</body>", "").strip()
    answers_path = os.path.join(out_dir, stem + "-answers.html")
    with open(answers_path, "w", encoding="utf-8") as f:
        f.write(head + ans_body + "\n</body>\n</html>\n")

    return student_path, answers_path


def convert_docx(html_path, docx_path):
    """调用 html_to_docx 插件转换 (与命令行 -m html_to_docx convert 等价)。"""
    scripts_dir = find_html_to_docx()
    env = dict(os.environ)
    env["PYTHONPATH"] = os.pathsep.join([scripts_dir, HTMLDOCX_LIBS])
    result = subprocess.run(
        [sys.executable, "-m", "html_to_docx", "convert", html_path,
         "-o", docx_path, "--page-size", "A4"],
        capture_output=True, text=True, env=env)
    if result.returncode != 0:
        print(result.stdout[-800:])
        print(result.stderr[-800:])
        raise SystemExit("[ERROR] 转换失败: %s" % html_path)
    if not os.path.isfile(docx_path):
        raise SystemExit("[ERROR] 未生成文件: %s" % docx_path)


def verify_docx(path, expect_answer):
    """校验单个 docx: 全黑(含样式层) / 表格 / (答案版)含参考答案字样。返回问题列表。"""
    sys.path.insert(0, HTMLDOCX_LIBS)
    import docx  # noqa: E402  (位于 htmldocx-libs 环境)
    import collections  # noqa: E402

    problems = []
    d = docx.Document(path)
    colors = collections.Counter()
    style_color_hits = collections.Counter()
    for p in d.paragraphs:
        style_color = _style_font_color(p)
        for r in p.runs:
            c = r.font.color.rgb if (
                r.font.color and r.font.color.type is not None
                and r.font.color.rgb) else None
            colors[str(c)] += 1
            # 有效颜色 = run 色或段落样式色; 样式色非黑即漏网(如 Heading 默认蓝)
            if c is None and style_color is not None and str(style_color) != "000000":
                style_color_hits["%s->%s" % (p.style.name, style_color)] += 1
    bad = [c for c in colors if c not in ("None", "000000")]
    if bad:
        problems.append("存在非黑色 run 颜色: %s" % bad)
    if style_color_hits:
        problems.append("标题继承非黑样式色(需 force_black): %s" % dict(style_color_hits))
    has_answer_text = any("参考答案" in p.text for p in d.paragraphs)
    if expect_answer and not has_answer_text:
        problems.append("答案版缺少'参考答案'字样")
    if not expect_answer and has_answer_text:
        problems.append("学生版出现'参考答案'字样!")
    n_tables = len(d.tables)
    print("  [校验] %s | run颜色: %s | 样式色问题: %s | 表格数: %d | 含答案: %s" % (
        os.path.basename(path), dict(colors), dict(style_color_hits) or "无",
        n_tables, has_answer_text))
    return problems


def _style_font_color(p):
    """取段落样式上显式定义的字体颜色, 无则 None。"""
    try:
        if (p.style is not None and p.style.font is not None
                and p.style.font.color is not None
                and p.style.font.color.type is not None):
            return p.style.font.color.rgb
    except Exception:
        pass
    return None


def force_black(path):
    """把所有'有效颜色非黑'的 run 强制设为纯黑。

    针对 html-to-docx 将 h1/h2 映射到 Word 内置 Heading 样式(默认蓝 365F91/4F81BD)、
    而 CSS body 颜色无法穿透样式层的问题。run 显式设黑色后覆盖样式色。
    """
    sys.path.insert(0, HTMLDOCX_LIBS)
    import docx  # noqa: E402
    from docx.shared import RGBColor  # noqa: E402

    d = docx.Document(path)
    changed = 0
    paras = list(d.paragraphs)
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                paras.extend(cell.paragraphs)
    for p in paras:
        style_color = _style_font_color(p)
        style_is_bad = style_color is not None and str(style_color) != "000000"
        for r in p.runs:
            rc = r.font.color.rgb if (
                r.font.color and r.font.color.type is not None
                and r.font.color.rgb) else None
            if style_is_bad or (rc is not None and str(rc) != "000000"):
                r.font.color.rgb = RGBColor(0, 0, 0)
                changed += 1
    if changed:
        d.save(path)
    return changed


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", required=True, help="合一版试卷 HTML 路径")
    ap.add_argument("--out", required=True, help="输出目录")
    ap.add_argument("--title", required=True, help="试卷名, 如 九上道法第一单元测试卷")
    args = ap.parse_args()

    src_path = os.path.abspath(args.src)
    out_dir = os.path.abspath(args.out)
    os.makedirs(out_dir, exist_ok=True)
    stem = os.path.join(out_dir, re.sub(r'[\\/:*?"<>|]', "", args.title))

    print("[1/3] 拆分学生版 / 答案版 ...")
    student_html, answers_html = split_html(src_path, out_dir, stem)
    print("  学生版 HTML: %s" % student_html)
    print("  答案版 HTML: %s" % answers_html)

    print("[2/3] 转换 DOCX ...")
    student_docx = stem + "-学生版.docx"
    answers_docx = stem + "-答案版.docx"
    convert_docx(student_html, student_docx)
    convert_docx(answers_html, answers_docx)
    print("  %s" % student_docx)
    print("  %s" % answers_docx)

    print("[2.5/3] 强制标题黑色 (修复 Word Heading 样式默认蓝) ...")
    n1 = force_black(student_docx)
    n2 = force_black(answers_docx)
    print("  已强制黑色 run: 学生版 %d 处 / 答案版 %d 处" % (n1, n2))

    print("[3/3] 格式校验 ...")
    problems = []
    problems += verify_docx(student_docx, expect_answer=False)
    problems += verify_docx(answers_docx, expect_answer=True)
    if problems:
        print("[FAIL] 校验未通过:")
        for p in problems:
            print("  - " + p)
        sys.exit(1)
    print("[OK] 全部校验通过: 全黑字体 / 学生版无答案 / 答案版结构完整")


if __name__ == "__main__":
    main()
